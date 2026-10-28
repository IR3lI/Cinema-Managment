import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_demo_docx():
    doc = Document()
    
    # Page setup: A4, 2cm margins
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    navy = RGBColor(15, 23, 42)
    primary = RGBColor(5, 150, 105) # Emerald primary
    dark_gray = RGBColor(51, 65, 85)
    accent = RGBColor(225, 29, 72)
    blue_accent = RGBColor(29, 78, 216)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.25, bold=False, italic=False, size=12, color=dark_gray):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = primary
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = navy
        return p

    # Header / Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("KỊCH BẢN DEMO SẢN PHẨM TRỰC TIẾP (05 PHÚT)")
    run_t.font.name = "Times New Roman"
    run_t.font.size = Pt(18)
    run_t.font.bold = True
    run_t.font.color.rgb = primary

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("HỆ THỐNG CINEMANAGEMENT - HƯỚNG DẪN THAO TÁC VÀ LỜI THOẠI CHO GIÁO VIÊN\nĐề tài: Quản lý dự án CNTT | GVHD: TS. Trần Khánh Dung")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = navy

    # Info box
    info_table = doc.add_table(rows=1, cols=1)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = info_table.rows[0].cells[0]
    cell._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="ECFDF5"/>'.format(nsdecls('w'))))
    set_cell_margins(cell, 120, 120, 200, 200)
    p_info = cell.paragraphs[0]
    p_info.paragraph_format.space_after = Pt(0)
    p_info.paragraph_format.line_spacing = 1.2
    
    r1 = p_info.add_run("• Thời lượng Demo: ")
    r1.font.bold = True
    p_info.add_run("Đúng 05 phút (300 giây) | Thao tác mượt mà theo kịch bản\n")
    r2 = p_info.add_run("• Setup kỹ thuật: ")
    r2.font.bold = True
    p_info.add_run("Chia đôi màn hình (Split Screen): Trái (Chrome - Khách hàng), Phải (Edge - Nhân viên & Admin)\n")
    r3 = p_info.add_run("• Tài khoản thực nghiệm: ")
    r3.font.bold = True
    p_info.add_run("Customer (customer@cinema.com / Customer@123) | Staff (staff@cinema.com / Staff@123) | Admin (admin@cinema.com / Admin@123)")
    for run in p_info.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

    add_p("", space_after=8)

    # Section 1: Timing Sheet
    add_h1("I. BẢNG PHÂN BỔ THỜI GIAN DEMO (TIMING SHEET - ĐÚNG 300 GIÂY)")
    
    table_data = [
        ["Thời gian", "Phân đoạn", "Vai trò & Màn hình", "Thao tác chính", "Điểm kỹ thuật mấu chốt"],
        ["0:00 - 0:35", "1. Trang chủ", "Khách hàng (Chrome)", "Lướt trang chủ, banner, poster, ngày tiếng Việt", "Logo chuẩn, UI hiện đại"],
        ["0:35 - 1:50", "2. Đặt vé trọn gói", "Khách hàng (Chrome)", "Chọn phim -> Chọn 2 ghế -> Chọn F&B -> VietQR", "Khóa ghế 10p, gộp 1 mã QR, gửi Gmail"],
        ["1:50 - 3:10", "3. Soát vé QR", "Nhân viên Staff (Edge)", "Đăng nhập Staff -> Quét Camera -> Quét lại lần 2", "Quét < 2s, quét lần 2 báo đỏ cảnh báo"],
        ["3:10 - 4:25", "4. Dashboard & BI", "Quản trị Admin (Edge)", "Đăng nhập Admin -> Xem Báo cáo & Thống kê", "Doanh thu phim cao nhất, Combo theo ngày"],
        ["4:25 - 5:00", "5. Kết luận & EXE", "Hội đồng", "Mở thư mục file .exe độc lập -> Mời Thầy/Cô Q&A", "Chạy không cần SDK .NET, Worker 30s"]
    ]

    t = doc.add_table(rows=len(table_data), cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.0), Inches(1.1), Inches(1.5), Inches(1.8), Inches(1.8)]
    
    for r_idx, row in enumerate(t.rows):
        is_hdr = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            cell.width = widths[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            if is_hdr:
                tcPr.append(parse_xml(r'<w:shd {} w:fill="065F46"/>'.format(nsdecls('w'))))
            elif r_idx % 2 == 1:
                tcPr.append(parse_xml(r'<w:shd {} w:fill="F8FAFC"/>'.format(nsdecls('w'))))
            set_cell_margins(cell, 80, 80, 80, 80)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            txt = table_data[r_idx][c_idx]
            run = p.add_run(txt)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10 if not is_hdr else 10.5)
            run.font.bold = is_hdr or (c_idx == 0)
            if is_hdr:
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = primary if c_idx == 0 else dark_gray

    add_p("", space_after=10)

    # Section 2: Step-by-step
    add_h1("II. KỊCH BẢN THAO TÁC VÀ LỜI THOẠI CHI TIẾT (STEP-BY-STEP LIVE SCRIPT)")

    demo_steps = [
        ("PHÂN ĐOẠN 1: KHỞI ĐỘNG & TỔNG QUAN KHÁCH HÀNG (0:00 - 0:35)",
         "Màn hình: Trình duyệt Trái (Chrome) - Trang chủ http://localhost:5006",
         "1. Cuộn nhẹ chuột từ trên xuống dưới trang chủ: lướt qua Hero Banner, danh sách phim Đang chiếu, Đánh giá sao của khán giả.\n2. Chỉ chuột vào thanh menu điều hướng và logo CineManagement.",
         "\"Em kính chào Thầy/Cô và Hội đồng! Tiếp theo phần thuyết trình, em xin phép được demo trực tiếp toàn bộ luồng vận hành của hệ thống CineManagement.\n"
         "Đầu tiên, đây là giao diện trang chủ dành cho Khách hàng. Hệ thống đã được chuẩn hóa nhận diện thương hiệu CineManagement, hỗ trợ xem danh sách phim đang chiếu, sắp chiếu, trailer và đánh giá sao thực tế từ người xem. Toàn bộ lịch chiếu và thời gian đều được bản địa hóa 100% bằng tiếng Việt.\""),

        ("PHÂN ĐOẠN 2: TRẢI NGHIỆM ĐẶT VÉ, KHÓA GHẾ 10P & VIETQR (0:35 - 1:50)",
         "Màn hình: Trình duyệt Trái (Chrome)",
         "1. Bấm 'Đặt vé ngay' một phim đang chiếu (ví dụ: Mai hoặc Đào, Phở và Piano).\n"
         "2. Chọn ngày chiếu hôm nay và chọn một suất chiếu cụ thể.\n"
         "3. Tại sơ đồ phòng chiếu: Click chọn 2 ghế cạnh nhau (F5, F6) -> Chỉ vào đồng hồ đếm ngược 09:59 và 2 ghế chuyển sang màu cam/vàng (Held).\n"
         "4. Bấm 'Tiếp tục chọn Bắp Nước' -> Chọn 1 Combo Bắp Nước Siêu Tiết Kiệm (có ảnh minh họa).\n"
         "5. Bấm 'Tiến hành Thanh toán' -> Hiển thị cổng VietQR Napas 247 -> Bấm 'Xác nhận đã thanh toán'.\n"
         "6. Tại trang Xác nhận thành công: Chỉ rõ hệ thống đã GỘP 2 GHẾ VÀO 1 MÃ QR DUY NHẤT và thông báo vé đã gửi về Gmail.",
         "\"Bây giờ em sẽ thực hiện đặt vé cho khách hàng. Em chọn phim, chọn suất chiếu. Tại sơ đồ phòng chiếu, em chọn 2 ghế F5 và F6.\n"
         "Thưa cô, điểm kỹ thuật đặc biệt ở đây là ngay khi em chọn ghế, cơ chế Seat Hold Lock lập tức kích hoạt: 2 ghế này chuyển sang trạng thái giữ chỗ với đồng hồ đếm ngược đúng 10 phút. Nếu có một khách hàng khác ở máy tính khác mở lên lúc này thì 2 ghế F5, F6 sẽ bị khóa, hoàn toàn không thể tranh chấp, loại bỏ 100% lỗi Race Condition.\n"
         "Tiếp theo, khách hàng chọn thêm Combo Bắp nước có hình ảnh trực quan. Tại trang thanh toán, hệ thống tạo mã VietQR tiêu chuẩn quốc gia, khách hàng chỉ cần quét app ngân hàng là tiền về đúng số tài khoản.\n"
         "Sau khi thanh toán thành công, theo đúng yêu cầu thực tế của rạp, hệ thống tự động gộp cả 2 ghế vào 1 mã QR duy nhất, đồng thời kích hoạt Background Worker gửi email vé điện tử HTML về Gmail của khách hàng.\""),

        ("PHÂN ĐOẠN 3: NHÂN VIÊN SOÁT VÉ BẰNG CAMERA & CHỐNG VÉ GIA LẬN (1:50 - 3:10)",
         "Màn hình: Trình duyệt Phải (Edge)",
         "1. Đăng nhập tài khoản Staff: staff@cinema.com / Mật khẩu: Staff@123.\n"
         "2. Bấm Avatar góc phải -> mở 'Hồ sơ cá nhân' (Profile nhân viên mới cập nhật thẻ nhân viên, ca trực).\n"
         "3. Bấm nút màu xanh 'Soát vé QR' trên Menu (/Admin/ScanTicket) -> Bấm 'Bật Camera'.\n"
         "4. Quét lần 1 (Vé hợp lệ): Đưa mã QR trên màn hình điện thoại hoặc ảnh QR vào camera -> Màn hình lóe xanh, phát âm thanh bíp và thông báo 'Check-in thành công! Suất chiếu: ... Ghế: F5, F6'.\n"
         "5. Quét lần 2 (Kiểm thử vé đã dùng): Đưa lại mã QR đó quét lần nữa -> Màn hình hiện khung cảnh báo đỏ 'CẢNH BÁO: Vé này đã được check-in vào lúc [HH:mm] bởi nhân viên Staff! Từ chối vào rạp!'.",
         "\"Bây giờ em chuyển sang vai trò Nhân viên soát vé tại cửa rạp. Em đăng nhập bằng tài khoản staff@cinema.com.\n"
         "Đây là hồ sơ cá nhân của nhân viên với đầy đủ thông tin ca trực và mã số. Em bấm vào tính năng 'Soát vé QR'. Ứng dụng tích hợp thư viện html5-qrcode, cho phép sử dụng trực tiếp Camera laptop hoặc Camera điện thoại mà không cần mua đầu đọc đắt tiền.\n"
         "Em bật Camera và đưa mã QR đơn hàng vừa đặt vào... Thưa cô, trong chưa đầy 1.5 giây, hệ thống đã xác thực thành công cả 2 ghế F5, F6 và đổi trạng thái vé sang Đã sử dụng.\n"
         "Và để kiểm tra tính năng bảo mật: Giả sử khách hàng chụp lại màn hình mã QR này gửi cho bạn bè để vào rạp lần thứ hai. Em đưa lại mã này trước camera... Hệ thống lập tức hiện cảnh báo đỏ: Vé đã được sử dụng lúc [giờ:phút], ngăn chặn tuyệt đối tình trạng gian lận vé!\""),

        ("PHÂN ĐOẠN 4: ADMIN DASHBOARD & BÁO CÁO DOANH THU CHUYÊN SÂU (3:10 - 4:25)",
         "Màn hình: Trình duyệt Phải (Edge)",
         "1. Đăng xuất Staff -> Đăng nhập tài khoản Quản trị: admin@cinema.com / Mật khẩu: Admin@123.\n"
         "2. Vào Admin Dashboard: Chỉ vào các thẻ KPI (doanh thu và vé hôm nay vừa tự động nhảy số theo đơn vừa đặt).\n"
         "3. Bấm 'Xem báo cáo' sang trang Báo cáo & Thống kê (/Admin/Reports):\n"
         "   - Chỉ vào Mục 1: Bảng xếp hạng phim có doanh thu cao nhất, số vé và tỷ lệ lấp đầy.\n"
         "   - Chỉ vào Mục 2: Biểu đồ doanh thu combo F&B theo từng ngày.\n"
         "   - Bấm thử bộ lọc: 'Hôm nay', '7 ngày qua', 'Tháng này'.\n"
         "4. Vào mục Lịch chiếu: Thử thêm 1 suất chiếu trùng giờ và trùng phòng -> Hệ thống báo lỗi chặn trùng lịch.",
         "\"Tiếp theo, em xin phép demo phân hệ Quản trị viên cấp cao. Em đăng nhập tài khoản admin@cinema.com.\n"
         "Tại Dashboard quản trị, các chỉ số doanh thu và số vé hôm nay vừa được cập nhật tức thời theo thời gian thực ngay sau đơn đặt vé vừa rồi.\n"
         "Em bấm vào mục Báo cáo & Thống kê Doanh thu chi tiết: Hệ thống phân tích sâu 2 bài toán kinh doanh cốt lõi của rạp:\n"
         "Một là Bảng xếp hạng phim có doanh thu cao nhất: Giúp quản lý rạp biết phim nào đang ăn khách để tăng thêm suất chiếu.\n"
         "Hai là Báo cáo doanh thu đồ ăn nước uống F&B theo từng ngày: Giúp kiểm soát tỷ trọng bán bắp nước, kiểm kê kho và đánh giá hiệu quả của các gói combo.\n"
         "Ngoài ra, tại phần Quản lý suất chiếu, hệ thống có thuật toán kiểm tra xung đột thời gian, tự động ngăn chặn hoàn toàn việc nhân viên xếp trùng giờ chiếu trong cùng một phòng.\""),

        ("PHÂN ĐOẠN 5: ĐÓNG GÓI EXE ĐỘC LẬP & TỔNG KẾT (4:25 - 5:00)",
         "Màn hình: Mở cửa sổ File Explorer / Terminal",
         "1. Chỉ vào file CineManagement.exe độc lập và file database app.db (SQLite portable).\n"
         "2. Chuyển ánh mắt về phía Thầy/Cô, nở nụ cười và phát biểu kết thúc.",
         "\"Điểm cuối cùng nhóm muốn chia sẻ là toàn bộ hệ thống đã được tối ưu và đóng gói thành một file EXE Portable duy nhất. Ứng dụng có thể copy sang bất kỳ máy tính Windows nào và khởi chạy ngay lập tức mà không cần cài đặt .NET SDK hay máy chủ MySQL phức tạp.\n"
         "Hệ thống cũng tích hợp một Background Service chạy ngầm định kỳ 30 giây để tự động quét dọn và giải phóng các ghế giữ chỗ quá 10 phút mà không làm tiêu tốn tài nguyên máy.\n"
         "Đến đây, nhóm em đã hoàn thành trọn vẹn kịch bản demo sản phẩm theo đúng cam kết trong đồ án. Nhóm em xin chân thành cảm ơn Thầy/Cô và kính mời Thầy/Cô đưa ra các câu hỏi nhận xét hoặc yêu cầu nhóm thao tác thêm bất kỳ chức năng nào ạ!\"")
    ]

    for title, screen_info, steps_info, spoken_info in demo_steps:
        add_h1(title)
        
        # Step action box
        step_table = doc.add_table(rows=1, cols=1)
        step_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_step = step_table.rows[0].cells[0]
        c_step._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w'))))
        set_cell_margins(c_step, 80, 80, 120, 120)
        p_s = c_step.paragraphs[0]
        p_s.paragraph_format.space_after = Pt(2)
        r_scr = p_s.add_run(f"🖥️ {screen_info}\n")
        r_scr.font.bold = True
        r_scr.font.color.rgb = blue_accent
        
        r_step_lbl = p_s.add_run("📋 Các thao tác cần bấm:\n")
        r_step_lbl.font.bold = True
        p_s.add_run(steps_info)
        for r in p_s.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

        add_p("", space_after=4)

        # Spoken box
        add_h2("🗣️ Lời nói trực tiếp với Giáo viên:")
        spk_table = doc.add_table(rows=1, cols=1)
        spk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_spk = spk_table.rows[0].cells[0]
        c_spk._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="ECFDF5"/>'.format(nsdecls('w'))))
        set_cell_margins(c_spk, 100, 100, 150, 150)
        
        p_spk = c_spk.paragraphs[0]
        p_spk.paragraph_format.space_after = Pt(0)
        p_spk.paragraph_format.line_spacing = 1.25
        
        for idx, line in enumerate(spoken_info.split('\n')):
            if idx > 0:
                p_spk = c_spk.add_paragraph()
                p_spk.paragraph_format.space_after = Pt(0)
                p_spk.paragraph_format.line_spacing = 1.25
            r = p_spk.add_run(line.strip())
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.5)
            r.font.italic = True
            r.font.color.rgb = navy

        add_p("", space_after=10)

    # Section 3: Troubleshooting
    add_h1("III. BẢNG DỰ PHÒNG TÌNH HUỐNG (TROUBLESHOOTING KHI DEMO)")

    tb_data = [
        ["Tình huống tại phòng thi", "Nguyên nhân", "Cách xử lý nhanh trong 3 giây"],
        ["Trình duyệt không nhận Camera", "Máy chiếu / Quyền camera bị chặn", "Bấm ngay 'Tải/Chụp ảnh QR' hoặc click Vé mẫu ở bảng bên dưới. Kết quả vẫn kiểm tra vé 100%!"],
        ["Bấm thanh toán không có tiền thật", "Chế độ demo học thuật", "Bấm 'Xác nhận đã thanh toán', hệ thống tự động hoàn tất và sinh mã QR."],
        ["Cô giáo hỏi kiểm chứng 10 phút nhả ghế", "Cô nghi ngờ tính năng chạy ngầm", "Mở file SeatCleanupWorker.cs cho cô xem BackgroundService chu kỳ 30s với câu truy vấn HoldExpiresAt."]
    ]

    t_tb = doc.add_table(rows=len(tb_data), cols=3)
    t_tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths_tb = [Inches(2.2), Inches(2.2), Inches(2.8)]

    for r_idx, row in enumerate(t_tb.rows):
        is_hdr = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            cell.width = widths_tb[c_idx]
            tcPr = cell._tc.get_or_add_tcPr()
            if is_hdr:
                tcPr.append(parse_xml(r'<w:shd {} w:fill="1E293B"/>'.format(nsdecls('w'))))
            elif r_idx % 2 == 1:
                tcPr.append(parse_xml(r'<w:shd {} w:fill="FEF2F2"/>'.format(nsdecls('w'))))
            set_cell_margins(cell, 80, 80, 100, 100)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            txt = tb_data[r_idx][c_idx]
            run = p.add_run(txt)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            run.font.bold = is_hdr or (c_idx == 0)
            if is_hdr:
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                run.font.color.rgb = accent if c_idx == 0 else dark_gray

    out_file = "Kich_Ban_Demo_5_Phut_CineManagement.docx"
    doc.save(out_file)
    print(f"SUCCESS: Saved '{out_file}' ({os.path.getsize(out_file)} bytes)")

if __name__ == "__main__":
    create_demo_docx()
