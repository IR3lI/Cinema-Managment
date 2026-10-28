import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1 pt = 20 twips)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_border(cell, **kwargs):
    """
    Set cell borders: top, bottom, left, right.
    val: 'single', 'nil', etc. color: 'HEX', sz: '4' (0.5 pt)
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_el = OxmlElement(f'w:{edge}')
            b_el.set(qn('w:val'), edge_data.get('val', 'single'))
            b_el.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), edge_data.get('color', 'CCCCCC'))
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def add_header_cell(row, col_idx, text, width_cm=None):
    cell = row.cells[col_idx]
    if width_cm:
        cell.width = Cm(width_cm)
    set_cell_shading(cell, "1E3A8A") # Navy Blue
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    set_cell_border(cell, 
                    top={'val': 'single', 'sz': 6, 'color': '1E3A8A'},
                    bottom={'val': 'single', 'sz': 8, 'color': '1E3A8A'},
                    left={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                    right={'val': 'single', 'sz': 4, 'color': 'CBD5E1'})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    return cell

def add_body_cell(row, col_idx, text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, bg_hex=None, width_cm=None):
    cell = row.cells[col_idx]
    if width_cm:
        cell.width = Cm(width_cm)
    if bg_hex:
        set_cell_shading(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_border(cell, 
                    top={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                    bottom={'val': 'single', 'sz': 4, 'color': 'CBD5E1'},
                    left={'val': 'single', 'sz': 4, 'color': 'E2E8F0'},
                    right={'val': 'single', 'sz': 4, 'color': 'E2E8F0'})
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(30, 41, 59)
    return cell

def add_image_placeholder(doc, fig_number, fig_caption):
    """Inserts a styled, professional placeholder box for images."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Cm(15.5)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=200, bottom=200, left=200, right=200)
    set_cell_border(cell,
                    top={'val': 'dashed', 'sz': 6, 'color': '94A3B8'},
                    bottom={'val': 'dashed', 'sz': 6, 'color': '94A3B8'},
                    left={'val': 'dashed', 'sz': 6, 'color': '94A3B8'},
                    right={'val': 'dashed', 'sz': 6, 'color': '94A3B8'})
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run1 = p.add_run("📷 [KHUNG DÁN ẢNH CHỤP MÀN HÌNH]\n")
    run1.font.name = "Times New Roman"
    run1.font.size = Pt(11)
    run1.font.bold = True
    run1.font.color.rgb = RGBColor(71, 85, 105)
    
    run2 = p.add_run(f"Dán ảnh cho {fig_caption} tại đây")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(10)
    run2.font.italic = True
    run2.font.color.rgb = RGBColor(100, 116, 139)

    # Caption below the placeholder
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(4)
    cp.paragraph_format.space_after = Pt(10)
    c_run = cp.add_run(f"Hình {fig_number}: {fig_caption}")
    c_run.font.name = "Times New Roman"
    c_run.font.size = Pt(10.5)
    c_run.font.italic = True
    c_run.font.bold = True
    c_run.font.color.rgb = RGBColor(30, 41, 59)

def add_diagram_image(doc, fig_number, fig_caption, image_path, width_inches=6.2):
    """Inserts a real image diagram if file exists, otherwise falls back to styled placeholder."""
    if os.path.exists(image_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(image_path, width=Inches(width_inches))
        
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(10)
        c_run = cp.add_run(f"Hình {fig_number}: {fig_caption}")
        c_run.font.name = "Times New Roman"
        c_run.font.size = Pt(10.5)
        c_run.font.italic = True
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(30, 41, 59)
    else:
        add_image_placeholder(doc, fig_number, fig_caption)

print("Helper functions ready.")
